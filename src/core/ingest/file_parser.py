"""Streaming server-side file parsers for the async ingest pipeline.

Each parser yields one dict per event row without loading the full file into
memory. This is the critical difference from the browser-side parsers in
breach.js which call JSON.parse(entire file text) on the main thread.

Supported formats: CSV, NDJSON, JSON (top-level array or well-known keys),
XLSX (read-only streaming via openpyxl).

ijson is used for large JSON files; if unavailable the parser falls back to
stdlib json with a warning.
"""
from __future__ import annotations

import csv
import json
import logging
import os
from typing import Iterator

logger = logging.getLogger(__name__)

# Well-known top-level array keys from telemetry exports
_JSON_ARRAY_KEYS = (
    "events", "logs", "records", "data", "items", "results",
    "Records",          # AWS CloudTrail
    "value",            # MS Graph / Defender
    "hits",             # Elasticsearch _search
)


def parse_file(path: str, filename: str | None = None) -> Iterator[dict]:
    """Detect file type and stream rows, yielding one dict per event.

    Injects ``_source`` (filename) into every row so downstream normalizers
    can identify the originating file.
    """
    name = filename or os.path.basename(path)
    ext = os.path.splitext(name)[1].lower()
    if ext in (".ndjson", ".jsonl"):
        yield from _parse_ndjson(path, name)
    elif ext == ".csv":
        yield from _parse_csv(path, name)
    elif ext in (".xlsx", ".xlsm"):
        yield from _parse_xlsx(path, name)
    elif ext in (".json",):
        yield from _parse_json(path, name)
    else:
        # Try NDJSON first (common for custom log exports), then JSON
        try:
            count = 0
            for row in _parse_ndjson(path, name):
                count += 1
                yield row
            if count == 0:
                yield from _parse_json(path, name)
        except Exception:
            yield from _parse_json(path, name)


# ── NDJSON ────────────────────────────────────────────────────────────────────

def _parse_ndjson(path: str, source: str) -> Iterator[dict]:
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
                if isinstance(obj, dict):
                    obj.setdefault("_source", source)
                    yield obj
                elif isinstance(obj, list):
                    for item in obj:
                        if isinstance(item, dict):
                            item.setdefault("_source", source)
                            yield item
            except json.JSONDecodeError:
                pass


# ── CSV ───────────────────────────────────────────────────────────────────────

def _parse_csv(path: str, source: str) -> Iterator[dict]:
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            d = dict(row)
            d["_source"] = source
            yield d


# ── JSON ──────────────────────────────────────────────────────────────────────

def _parse_json(path: str, source: str) -> Iterator[dict]:
    """Stream a JSON file. Uses ijson for large files, stdlib for small ones."""
    size = os.path.getsize(path)
    if size > 5 * 1024 * 1024:  # >5 MB — use ijson streaming
        yield from _parse_json_ijson(path, source)
    else:
        yield from _parse_json_stdlib(path, source)


def _parse_json_stdlib(path: str, source: str) -> Iterator[dict]:
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as fh:
            data = json.load(fh)
    except Exception as exc:
        logger.warning("JSON parse failed for %s: %s", path, exc)
        return
    yield from _flatten_json_data(data, source)


def _parse_json_ijson(path: str, source: str) -> Iterator[dict]:
    try:
        import ijson
    except ImportError:
        logger.warning("ijson not installed — falling back to stdlib JSON for %s", path)
        yield from _parse_json_stdlib(path, source)
        return

    # Stream every top-level array. This handles telemetry bundles such as
    # {aws_cloudtrail: [...], okta_system_log: [...]} without materialising
    # the arrays with list(...).
    emitted = False
    try:
        with open(path, "rb") as fh:
            for prefix, event, value in ijson.parse(fh):
                if prefix and "." not in prefix and event == "start_array":
                    with open(path, "rb") as rows_fh:
                        for item in ijson.items(rows_fh, f"{prefix}.item"):
                            if isinstance(item, dict):
                                item.setdefault("_source", source)
                                item.setdefault("_section", prefix)
                                emitted = True
                                yield item
        if emitted:
            return
    except Exception as exc:
        logger.debug("ijson top-level array scan failed for %s: %s", path, exc)

    # Fallback: stream top-level JSON arrays via item prefix.
    try:
        with open(path, "rb") as fh:
            for item in ijson.items(fh, "item"):
                if isinstance(item, dict):
                    item.setdefault("_source", source)
                    yield item
        return
    except Exception:
        pass

    # Last resort: parse whole file
    yield from _parse_json_stdlib(path, source)


def _flatten_json_data(data, source: str) -> Iterator[dict]:
    if isinstance(data, list):
        for item in data:
            if isinstance(item, dict):
                item.setdefault("_source", source)
                yield item
    elif isinstance(data, dict):
        # Try well-known array keys
        for key in _JSON_ARRAY_KEYS:
            if key in data and isinstance(data[key], list):
                for item in data[key]:
                    if isinstance(item, dict):
                        item.setdefault("_source", source)
                        yield item
                return
        # Multi-section object: flatten all top-level arrays
        found = False
        for v in data.values():
            if isinstance(v, list):
                for item in v:
                    if isinstance(item, dict):
                        item.setdefault("_source", source)
                        yield item
                        found = True
        if not found:
            data.setdefault("_source", source)
            yield data


# ── XLSX ──────────────────────────────────────────────────────────────────────

def _parse_xlsx(path: str, source: str) -> Iterator[dict]:
    """Stream an XLSX workbook in read-only mode — never loads entire workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.error("openpyxl not installed — cannot parse %s", path)
        return

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("XLSX open failed for %s: %s", path, exc)
        return

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers: list[str] | None = None
        for row in ws.iter_rows(values_only=True):
            if headers is None:
                headers = [str(c or "").strip() or f"col_{i}" for i, c in enumerate(row)]
                continue
            obj: dict = {"_source": source, "_sheet": sheet_name}
            for i, cell in enumerate(row):
                key = headers[i] if i < len(headers) else f"col_{i}"
                obj[key] = cell
            # Skip entirely empty rows
            if all(v is None or v == "" for k, v in obj.items() if not k.startswith("_")):
                continue
            yield obj
    wb.close()
