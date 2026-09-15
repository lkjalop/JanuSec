"""Input lane classification and evidence policy for uploaded assessment files.

The assessment pipeline must not treat context packs or evaluation answer keys
as telemetry evidence.  This module keeps that decision close to ingestion so
downstream clustering only sees rows that are valid evidence.

Evidence lanes
--------------
LANE_TELEMETRY_EVIDENCE   — uploaded logs/events; the only lane that can create findings.
LANE_BUSINESS_CONTEXT     — CMDB, HR, travel; may enrich/downgrade but cannot create findings.
LANE_EVALUATION_ANSWER_KEY — scoring packs, answer keys; offline use only, never production.
LANE_PLATFORM_PRIOR       — ATT&CK mappings, reputation heuristics; score deltas only.
LANE_LLM_NARRATIVE        — generated prose; terminal output only, never analysis input.
LANE_CACHE_OR_HISTORY     — prior assessments; display only, cannot influence current verdict.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

# ── Lane constants ─────────────────────────────────────────────────────────────
LANE_TELEMETRY_EVIDENCE = "telemetry_evidence"
LANE_BUSINESS_CONTEXT = "business_context"
LANE_EVALUATION_ANSWER_KEY = "evaluation_answer_key"
LANE_PLATFORM_PRIOR = "platform_prior"
LANE_LLM_NARRATIVE = "llm_narrative"
LANE_CACHE_OR_HISTORY = "cache_or_history"

# Lanes that are permitted to create confirmed/suspicious findings.
_FINDING_ALLOWED_LANES: frozenset[str] = frozenset({LANE_TELEMETRY_EVIDENCE})

# Attribute keys checked when validating a context downgrade.
_DOWNGRADE_ATTRIBUTE_KEYS = ("entity", "time_window", "source_infra", "destination_infra", "tool", "account", "approved_scope")

TELEMETRY_SHEETS = {
    "identity",
    "endpoint",
    "network",
    "cloud_aws",
    "cloud_azure",
    "email",
    "data_movement",
    "change_context",
}

EVALUATION_ANSWER_KEY_SHEETS = {
    "hr_directory",
    "asset_inventory",
    "vendors",
    "approved_travel",
    "threat_intel_iocs",
    "red_herring_sow",
    "crown_jewels_register",
}

BUSINESS_CONTEXT_SHEETS = {
    "hr_directory",
    "asset_inventory",
    "vendors",
    "approved_travel",
    "change_outage_calendar",
}


def _canon(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


def is_non_evidence_sheet(sheet_name: Any) -> bool:
    name = _canon(sheet_name)
    return name in EVALUATION_ANSWER_KEY_SHEETS or name in BUSINESS_CONTEXT_SHEETS


def classify_xlsx_sheets(sheet_names: Iterable[str], *, filename: str = "") -> dict[str, Any]:
    """Classify an XLSX workbook into an ingestion lane.

    Returns a small serialisable dict so API endpoints can expose the reason
    without sharing workbook contents.
    """
    names = {_canon(name) for name in sheet_names if str(name or "").strip()}
    telemetry_hits = sorted(names & TELEMETRY_SHEETS)
    answer_key_hits = sorted(names & EVALUATION_ANSWER_KEY_SHEETS)
    context_hits = sorted(names & BUSINESS_CONTEXT_SHEETS)

    if len(answer_key_hits) >= 3 or (
        "threat_intel_iocs" in names and "crown_jewels_register" in names
    ):
        return {
            "lane": "evaluation_answer_key",
            "evidence_allowed": False,
            "reason": "workbook matches evaluation/context answer-key sheet signature",
            "matched_sheets": answer_key_hits,
            "filename": filename,
        }

    if telemetry_hits:
        return {
            "lane": "telemetry_evidence",
            "evidence_allowed": True,
            "reason": "workbook contains recognised telemetry sheets",
            "matched_sheets": telemetry_hits,
            "filename": filename,
        }

    if context_hits:
        return {
            "lane": "business_context",
            "evidence_allowed": False,
            "reason": "workbook contains business context sheets but no telemetry sheets",
            "matched_sheets": context_hits,
            "filename": filename,
        }

    return {
        "lane": "telemetry_evidence",
        "evidence_allowed": True,
        "reason": "no context/answer-key signature detected",
        "matched_sheets": [],
        "filename": filename,
    }


def classify_xlsx_path(path: str | Path, *, filename: str = "") -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {
            "lane": "unknown",
            "evidence_allowed": False,
            "reason": "openpyxl unavailable",
            "matched_sheets": [],
            "filename": filename or Path(path).name,
        }

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return classify_xlsx_sheets(wb.sheetnames, filename=filename or Path(path).name)
    finally:
        wb.close()


# ── Evidence policy helpers ────────────────────────────────────────────────────

def can_create_finding(lane: str) -> bool:
    """Return True only for lanes that are permitted to create confirmed/suspicious findings."""
    return lane in _FINDING_ALLOWED_LANES


def require_telemetry_provenance(finding: dict[str, Any]) -> list[str]:
    """Return a list of policy violations for a finding that claims to be confirmed/suspicious.

    A violation is returned when:
    - there are no sources at all, or
    - none of the sources have lane == LANE_TELEMETRY_EVIDENCE with at least one row_ref.
    """
    classification = str(finding.get("classification") or "")
    if classification not in ("confirmed_threat", "suspicious_unconfirmed"):
        return []

    sources: list[dict] = finding.get("sources") or []
    if not sources:
        return [f"finding '{classification}' has no sources"]

    violations: list[str] = []
    has_telemetry = any(
        s.get("lane") == LANE_TELEMETRY_EVIDENCE and (s.get("row_refs") or [])
        for s in sources
    )
    if not has_telemetry:
        violations.append(
            f"finding '{classification}' has no telemetry_evidence source with row_refs"
        )
    return violations


def validate_context_downgrade(finding: dict[str, Any], context: dict[str, Any]) -> bool:
    """Return True only when the context matches at least two attributes from the finding.

    A single weak match (e.g. only 'there was a pentest this month') is not enough.
    """
    matches = 0
    for attr in _DOWNGRADE_ATTRIBUTE_KEYS:
        finding_val = str(finding.get(attr) or "").strip().lower()
        context_val = str(context.get(attr) or "").strip().lower()
        if finding_val and context_val and (finding_val in context_val or context_val in finding_val):
            matches += 1
        if matches >= 2:
            return True
    return False


def tag_row_lane(row: dict[str, Any], lane: str) -> dict[str, Any]:
    """Return a shallow copy of *row* with '_lane' set to *lane*."""
    tagged = dict(row)
    tagged["_lane"] = lane
    return tagged
