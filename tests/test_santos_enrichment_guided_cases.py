import csv
import json
from pathlib import Path

import openpyxl

from src.api.deep_analyze_endpoints import _hydrate_assessment_semantics


FIXTURE_DIR = Path("dump/test files")


def _load_workbook_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: list[dict] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h or "").strip() for h in rows[0]]
        for row in rows[1:]:
            item = {headers[i] or f"col_{i}": row[i] for i in range(min(len(headers), len(row)))}
            item["_sheet"] = ws.title
            item["_source"] = path.name
            out.append(item)
    return out


def _load_santos_sample() -> list[dict]:
    rows: list[dict] = []

    # Workbook context is the authority for pentest scope, travel, benign noise,
    # threat intel, and crown-jewel impact.
    rows.extend(_load_workbook_rows(FIXTURE_DIR / "janusec_enrichment_context_v1.xlsx"))

    # Add the high-signal telemetry rows needed to prove the story without
    # loading all 44K fixture rows into this unit test.
    terms = (
        "wei.zhao@allianceforwarding.com",
        "shipping_amendment",
        "ctflogger.dll",
        "comsvcs.dll",
        "OneDriveUpdate",
        "aaron.blackwood",
        "SnowflakeFedRole",
        "integration-runner",
        "Launch Privileged Container",
        "COPY INTO",
        "31.216.148.17",
        "mega.nz",
        "rclone.exe",
        "Red Herring",
        "Abdul Mohammadi",
        "Minecraft",
        "45.79.214.88",
    )
    lowered = tuple(t.lower() for t in terms)

    with (FIXTURE_DIR / "janusec_endpoint_k8s_v1.ndjson").open(encoding="utf-8") as fh:
        for line in fh:
            if any(t in line.lower() for t in lowered):
                item = json.loads(line)
                item["_source"] = "janusec_endpoint_k8s_v1.ndjson"
                rows.append(item)

    with (FIXTURE_DIR / "janusec_network_v1.csv").open(encoding="utf-8", newline="") as fh:
        for row in csv.DictReader(fh):
            blob = json.dumps(row)
            if any(t in blob.lower() for t in lowered):
                row["_source"] = "janusec_network_v1.csv"
                rows.append(row)

    data = json.load((FIXTURE_DIR / "janusec_cloud_identity_v1.json").open(encoding="utf-8"))
    for value in data.values():
        if not isinstance(value, list):
            continue
        for row in value:
            blob = json.dumps(row)
            if any(t in blob.lower() for t in lowered):
                row["_source"] = "janusec_cloud_identity_v1.json"
                rows.append(row)

    return rows


def test_santos_enrichment_cases_surface_expected_story():
    assessment = {
        "assessment_id": "santos-story-test",
        "rows": _load_santos_sample(),
        "org": "santos",
    }

    hydrated = _hydrate_assessment_semantics(assessment)
    clusters = hydrated["correlation_clusters"]
    by_id = {c["cluster_id"]: c for c in clusters}

    breach = by_id["case-unknown-2026-a-breach"]
    assert breach["verdict"] == "VALIDATED_BREACH"
    assert breach["severity"] == "critical"
    assert "pre-existing data breach" in breach["tier1_prefill"]["headline_subtitle"]
    assert "Rachel Nakamura" in breach["tier1_prefill"]["what_happened"]
    assert "Snowflake" in breach["tier1_prefill"]["what_happened"]
    assert len(breach["tier1_prefill"]["evidence_chain"]) >= 5

    assert by_id["case-red-herring-pentest"]["verdict"] == "BENIGN_EXPECTED"
    assert by_id["case-jacob-vpn-gaming"]["verdict"] == "BENIGN_EXPECTED"
    assert by_id["case-approved-apac-travel"]["verdict"] == "BENIGN_EXPECTED"

    # The newspaper-style front page should be the breach story, not a generic
    # cloud API or IP pivot.
    assert clusters[0]["cluster_id"] == "case-unknown-2026-a-breach"

