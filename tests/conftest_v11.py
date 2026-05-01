"""Shared fixtures for v1.1 dataset tests.

Provides pre-loaded rows from all three v1.1 test files so individual test
modules can import them without repeating file-load logic.
"""
from __future__ import annotations

import csv
import json
import os
from typing import Any, Dict, List

import pytest

_DUMP = os.path.join(os.path.dirname(__file__), "../dump/test files")


def _xlsx_to_rows(path: str) -> List[Dict[str, Any]]:
    """Load an xlsx with multiple sheets into a flat list of row dicts."""
    try:
        import openpyxl
    except ImportError:
        return []

    wb = openpyxl.load_workbook(path, read_only=True)
    all_rows: List[Dict] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_rows = list(ws.iter_rows(values_only=True))
        if not sheet_rows:
            continue
        headers = [str(h or "").strip() for h in sheet_rows[0]]
        for i, row in enumerate(sheet_rows[1:], start=1):
            d: Dict[str, Any] = {headers[j]: row[j] for j in range(min(len(headers), len(row)))}
            d["_sheet"] = sheet_name
            d["_row_number"] = i
            d["row_index"] = len(all_rows)
            # Normalise timestamp key
            for ts_key in ("date_utc", "timestamp_utc", "timestamp", "ts"):
                if d.get(ts_key):
                    d["ts"] = str(d[ts_key])
                    break
            all_rows.append(d)
    wb.close()
    return all_rows


def _csv_to_rows(path: str) -> List[Dict[str, Any]]:
    with open(path, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for i, r in enumerate(rows):
        r["row_index"] = i
        r.setdefault("ts", r.get("timestamp_utc") or r.get("timestamp") or "")
    return rows


def _json_to_rows(path: str) -> List[Dict[str, Any]]:
    with open(path, encoding="utf-8") as f:
        d = json.load(f)
    events = d.get("events", [])
    for i, e in enumerate(events):
        e["row_index"] = i
        e.setdefault("ts", e.get("timestamp_utc") or "")
        # normalise user field
        e.setdefault("user", e.get("user_principal_name") or "")
        e.setdefault("src_ip", e.get("source_ip") or "")
    return events


@pytest.fixture(scope="session")
def ep_rows():
    return _xlsx_to_rows(os.path.join(_DUMP, "janusec_ep_endpoint.v1.1.xlsx"))


@pytest.fixture(scope="session")
def net_rows():
    return _csv_to_rows(os.path.join(_DUMP, "janusec_net_c2_bgp.v1.1.csv"))


@pytest.fixture(scope="session")
def okta_rows():
    return _json_to_rows(os.path.join(_DUMP, "janusec_okta_m365_events.v1.1.json"))


@pytest.fixture(scope="session")
def all_rows(ep_rows, net_rows, okta_rows):
    """All 241 events from all three sources combined."""
    return ep_rows + net_rows + okta_rows


# ── Partition helpers ────────────────────────────────────────────────────────

def partition_by_state(rows, state_key="review_state"):
    malicious = [r for r in rows if str(r.get(state_key, "")).startswith("confirmed_malicious")]
    benign = [r for r in rows if str(r.get(state_key, "")).startswith("reviewed_benign")]
    needs = [r for r in rows if str(r.get(state_key, "")).startswith("needs_investigation")]
    script_kiddie = [r for r in rows if str(r.get(state_key, "")).startswith("script_kiddie")]
    return malicious, benign, needs, script_kiddie
