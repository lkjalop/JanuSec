import os, sys, json, asyncio
# make src resolvable
here = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
src_path = os.path.join(here, 'src')
if src_path not in sys.path:
    sys.path.insert(0, src_path)

try:
    from src.api.csv_handler import get_csv_processor
    from src.graph.hopgraph import GLOBAL_HOPGRAPH
except Exception:
    # ensure src is on sys.path
    if src_path not in sys.path:
        sys.path.insert(0, src_path)
    from api.csv_handler import get_csv_processor
    from graph.hopgraph import GLOBAL_HOPGRAPH

async def run():
    xlsx = os.path.join(here, 'dump', 'Cyberstash_csv2_sample.xlsx')
    with open(xlsx, 'rb') as f:
        raw = f.read()
    # emulate csv_endpoints Excel conversion path (CSVProcessor expects CSV bytes)
    # but the CSVProcessor provides process_csv which expects CSV bytes; csv_endpoints
    # does the conversion itself. To reuse, we'll duplicate simple conversion using openpyxl
    try:
        import io, csv
        from openpyxl import load_workbook
        bio = io.BytesIO(raw)
        wb = load_workbook(bio, read_only=True, data_only=True)
        sheet = wb.active
        buf = io.StringIO()
        writer = csv.writer(buf)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(['' if v is None else str(v) for v in row])
        content = buf.getvalue().encode('utf-8')
    except Exception as e:
        print('Excel conversion failed', e)
        return

    proc = get_csv_processor()
    res = await proc.process_csv(content, filename='Cyberstash_csv2_sample.xlsx')
    print(json.dumps(res, indent=2))

    # Ingest each processed result into HopGraph as an event mapping
    for r in res.get('results', []):
        raw_row = r.get('raw', {})
        event = {
            'timestamp': None,
            'host': raw_row.get('host'),
            'process': raw_row.get('process_name'),
            'pid': raw_row.get('pid'),
            'dst_ip': None,
            'src_ip': None,
            'domain': None,
            'file_hash': raw_row.get('hash'),
            'ja3': None,
            'cert_fp': None
        }
        try:
            from src.core.graph.hopgraph_utils import safe_upsert_node
        except Exception:
            safe_upsert_node = None
        try:
            if safe_upsert_node is not None and event.get('type') == 'file_hash' and event.get('id'):
                safe_upsert_node(GLOBAL_HOPGRAPH, 'file_hash', event.get('id'), attrs=event.get('attrs') or {}, source='csv_upload')
            else:
                GLOBAL_HOPGRAPH.ingest_event(event, source='csv_upload')
        except Exception:
            try:
                GLOBAL_HOPGRAPH.ingest_event(event, source='csv_upload')
            except Exception:
                pass

    print('HopGraph edges after ingest:', sum(len(v) for v in GLOBAL_HOPGRAPH.adj.values()))
    # Run explain_chain for host-alice
    ex = GLOBAL_HOPGRAPH.explain_chain(start='host:alice', max_depth=4, beam_width=6, top_k=5)
    print('Explain for host:alice after CSV ingest:')
    print(json.dumps(ex, indent=2))

if __name__ == '__main__':
    asyncio.run(run())
