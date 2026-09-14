from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from pathlib import Path
from typing import Any, Dict, List

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _load_manifest_rows(pack_dir: Path) -> List[Dict[str, Any]]:
    manifest_path = pack_dir / "manifest.json"
    if not manifest_path.exists():
        raise ValueError(f"pack directory missing manifest.json: {pack_dir}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    entries = manifest.get("files") or []
    rows: List[Dict[str, Any]] = []
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        rel = entry.get("path")
        if not rel:
            continue
        child_path = (pack_dir / rel).resolve()
        child_rows = _load_rows(child_path)
        source_kind = entry.get("source_kind")
        labels = entry.get("labels") or []
        for row in child_rows:
            row.setdefault("tenant_id", manifest.get("tenant_id"))
            row.setdefault("export_pack", pack_dir.name)
            row.setdefault("export_source", source_kind or Path(rel).stem)
            if labels:
                row.setdefault("pack_labels", list(labels))
        rows.extend(child_rows)
    return rows


def _load_xlsx_rows(path: Path) -> List[Dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows: List[Dict[str, Any]] = []
    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True))
        if not values:
            continue
        headers = [str(h).strip() if h is not None else f"col_{idx}" for idx, h in enumerate(values[0])]
        for idx, row in enumerate(values[1:], start=1):
            rec = {headers[i] or f"col_{i}": row[i] for i in range(min(len(headers), len(row)))}
            rec["row_index"] = len(rows)
            rec["sheet"] = ws.title
            rec["source_file"] = path.name
            rec["fingerprint"] = f"{path.name}:{ws.title}:{idx}"
            rows.append(rec)
    return rows


def _load_rows(path: Path) -> List[Dict[str, Any]]:
    if path.is_dir():
        return _load_manifest_rows(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _load_xlsx_rows(path)
    if suffix == ".csv":
        with path.open("r", encoding="utf-8", errors="ignore", newline="") as fh:
            reader = csv.DictReader(fh)
            rows = []
            for idx, row in enumerate(reader):
                rec = dict(row)
                rec["row_index"] = idx
                rec["sheet"] = path.stem
                rec["source_file"] = path.name
                rec["fingerprint"] = f"{path.name}:csv:{idx}"
                rows.append(rec)
            return rows
    if suffix in {".json", ".jsonl"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".jsonl":
            raw_rows = [json.loads(line) for line in text.splitlines() if line.strip()]
        else:
            payload = json.loads(text)
            raw_rows = payload if isinstance(payload, list) else payload.get("rows") or payload.get("records") or [payload]
        rows = []
        for idx, row in enumerate(raw_rows):
            rec = dict(row) if isinstance(row, dict) else {"value": row}
            rec["row_index"] = idx
            rec["sheet"] = path.stem
            rec["source_file"] = path.name
            rec["fingerprint"] = f"{path.name}:json:{idx}"
            rows.append(rec)
        return rows
    raise ValueError(f"unsupported input format: {path.suffix}")


def main() -> int:
    ap = argparse.ArgumentParser(description="Run offline replay against xlsx/csv/json/jsonl datasets.")
    ap.add_argument("paths", nargs="+", help="Dataset paths to replay")
    ap.add_argument("--tenant", default="offline-demo")
    ap.add_argument("--prod-scope", default="")
    ap.add_argument("--mock-llm", action="store_true")
    ap.add_argument("--out", default="")
    args = ap.parse_args()

    os.environ.setdefault("API_KEYS_JSON", '[{"key":"devkey123","scopes":["*"]}]')
    os.environ.setdefault("DEFAULT_FRONTEND", "console")
    os.environ["PROD_SCOPE"] = args.prod_scope
    os.environ["LLM_MOCK"] = "1" if args.mock_llm else os.environ.get("LLM_MOCK", "0")
    repo_root = Path(__file__).resolve().parents[1]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))

    from src.api.app import create_app

    client = TestClient(create_app())
    output: Dict[str, Any] = {"tenant": args.tenant, "prod_scope": args.prod_scope, "files": []}
    for raw_path in args.paths:
        path = Path(raw_path)
        rows = _load_rows(path)
        resp = client.post(
            "/api/v1/csv/deep_analyze",
            headers={"x-api-key": "devkey123", "X-Tenant-ID": args.tenant},
            json={"org": args.tenant, "rows": rows, "options": {"auto_llm": True}},
        )
        persona = client.post(
            "/api/v1/reports/persona_view?persona=executive&disclosure_level=2&top_n=5",
            headers={"x-api-key": "devkey123", "X-Tenant-ID": args.tenant},
            json=resp.json(),
        )
        tier2 = client.post(
            "/api/v1/csv/tier2_summarize",
            headers={"x-api-key": "devkey123", "X-Tenant-ID": args.tenant},
            json={"org": args.tenant, "assessment_id": resp.json().get("assessment_id"), "rows": rows[:25]},
        )
        output["files"].append(
            {
                "file": path.name,
                "path": str(path),
                "row_count": len(rows),
                "assessment": resp.json(),
                "persona": persona.json(),
                "tier2": tier2.json() if "application/json" in (tier2.headers.get("content-type") or "") else {"text": tier2.text},
            }
        )

    if args.out:
        Path(args.out).write_text(json.dumps(output, indent=2), encoding="utf-8")
    else:
        print(json.dumps(output, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
